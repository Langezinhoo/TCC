from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
SOURCE = ROOT / 'PrintWatch Artigo Final.docx'
OUTPUT = ROOT / 'Artigo PrintWatch Versao Final.docx'
ASSETS = ROOT / 'article-assets'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def configure_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, first_indent=1.25, line=1.5):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.first_line_indent = Cm(first_indent) if first_indent else None
    fmt.line_spacing = line


def add_body(doc, text):
    p = doc.add_paragraph(style='Normal')
    configure_paragraph(p)
    run = p.add_run(text)
    set_font(run)
    return p


def add_heading(doc, text, level=1):
    style = 'Heading 2' if level == 1 else 'Heading 3'
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(12 if level == 1 else 8)
    fmt.space_after = Pt(6)
    fmt.first_line_indent = None
    fmt.keep_with_next = True
    run = p.add_run(text)
    set_font(run, 12, bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_font(r, 10)
    return p


def add_figure(doc, filename, caption, width_cm=14.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    add_caption(doc, caption)
    add_caption(doc, 'Fonte: elaborado pelo autor (2026).')


def add_reference(doc, text):
    p = doc.add_paragraph(style='Normal')
    configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=4, first_indent=0, line=1.0)
    p.paragraph_format.left_indent = Cm(0)
    r = p.add_run(text)
    set_font(r, 10)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def setup_document(doc):
    clear_body(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    normal.font.size = Pt(12)
    for style_name in ('Title', 'Heading 2', 'Heading 3'):
        style = styles[style_name]
        style.font.name = 'Times New Roman'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer
    footer.paragraphs[0].text = ''
    footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_title_block(doc):
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run('DESENVOLVIMENTO DE UM SISTEMA WEB PARA MONITORAMENTO DE IMPRESSORAS EM REDE')
    set_font(r, 14, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('João Pedro Lange')
    set_font(r, 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run('Orientador: André de Freitas Smaira')
    set_font(r, 12)


def add_summary(doc):
    p = doc.add_paragraph()
    configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, first_indent=0, line=1.0)
    r = p.add_run('Resumo. ')
    set_font(r, 11, bold=True)
    r = p.add_run('Este trabalho apresenta o desenvolvimento do PrintWatch, um sistema web destinado ao monitoramento de impressoras em rede. A solução centraliza o cadastro de equipamentos, status operacional, contadores, níveis de toner, eventos técnicos, chamados de manutenção e controle de suprimentos. A aplicação foi construída com React no front-end, Node.js e Express no back-end e SQLite para persistência dos dados. Para demonstrar as funcionalidades sem depender de equipamentos físicos, foi criado um modo de dados simulados, que reproduz leituras de toner, contadores e situações operacionais. Os resultados mostram que o sistema oferece uma visão centralizada do parque de impressão, identifica situações que exigem atenção e registra informações que podem apoiar a manutenção preventiva. A integração automática por SNMP é indicada como evolução futura.')
    set_font(r, 11)
    p = doc.add_paragraph()
    configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=12, first_indent=0, line=1.0)
    r = p.add_run('Palavras-chave: ')
    set_font(r, 11, bold=True)
    r = p.add_run('monitoramento de impressoras; SNMP; manutenção preventiva; sistema web; gestão de suprimentos.')
    set_font(r, 11)


def build():
    doc = Document(SOURCE)
    setup_document(doc)
    add_title_block(doc)
    add_summary(doc)

    add_heading(doc, '1 INTRODUÇÃO')
    add_body(doc, 'O gerenciamento de dispositivos de impressão em ambientes corporativos exige acompanhamento contínuo de status, contadores de uso e suprimentos. Quando essas informações são verificadas manualmente, uma falha de comunicação, a falta de toner ou a indisponibilidade de um equipamento pode ser percebida apenas após afetar a rotina dos usuários. Esse cenário amplia o tempo de resposta do suporte e dificulta o planejamento de reposições e manutenções.')
    add_body(doc, 'O monitoramento centralizado permite registrar indicadores operacionais e apresentar informações de forma acessível para técnicos e gestores. Entre os dados relevantes estão o endereço IP, o estado de conectividade, o contador de páginas, os níveis de suprimentos e as ocorrências relacionadas ao equipamento. Esses elementos subsidiam a manutenção preventiva, pois permitem priorizar ações antes que uma interrupção se torne crítica.')
    add_body(doc, 'Este trabalho tem como objetivo desenvolver uma aplicação web para o monitoramento de impressoras em rede. O sistema proposto, denominado PrintWatch, reúne dados dos equipamentos, gera alertas operacionais, armazena histórico, registra chamados e apoia o controle de materiais. O artigo apresenta a fundamentação do problema, a arquitetura implementada, a modelagem dos dados e os resultados obtidos com a demonstração do sistema.')

    add_heading(doc, '2 FUNDAMENTAÇÃO TEÓRICA')
    add_heading(doc, '2.1 Monitoramento de dispositivos e manutenção preventiva', 2)
    add_body(doc, 'A gestão de impressão busca reduzir desperdícios, manter a disponibilidade dos equipamentos e organizar o atendimento técnico. O registro de contadores e ocorrências permite identificar padrões de uso e planejar a substituição de materiais. Para Kardec e Nascif (2013), a manutenção preventiva é orientada pela redução da probabilidade de falhas e pela preservação da disponibilidade dos ativos, princípio aplicável ao parque de impressão.')
    add_body(doc, 'Impressoras a laser e a jato de tinta possuem componentes e insumos com ciclos de consumo distintos. Por isso, o acompanhamento de toner, páginas impressas, falhas de comunicação e eventos técnicos facilita a priorização das atividades de suporte. Em ambientes com diversos setores ou unidades, uma interface centralizada reduz a dependência de verificações presenciais e torna as informações operacionais mais rastreáveis.')
    add_heading(doc, '2.2 Protocolo SNMP e aplicações web', 2)
    add_body(doc, 'O Simple Network Management Protocol (SNMP) é um protocolo de gerenciamento utilizado para consultar informações de dispositivos conectados à rede. A arquitetura do protocolo estabelece mecanismos para organizar dados de gerenciamento e permitir a comunicação entre agentes e sistemas gerenciadores (IETF, 2002). Em impressoras compatíveis, essa integração pode fornecer status, contadores e informações de suprimentos.')
    add_body(doc, 'Embora a coleta automática por SNMP não esteja ativada na versão demonstrada, o PrintWatch armazena parâmetros de conexão, como versão do protocolo e comunidade, deixando a estrutura preparada para essa evolução. A solução foi desenvolvida como uma aplicação web composta por interface em React, API REST em Node.js com Express e banco SQLite. Essa organização separa a apresentação, as regras de negócio e a persistência dos dados, facilitando a manutenção e futuras integrações (PRESSMAN; MAXIM, 2016).')

    add_heading(doc, '3 METODOLOGIA E DESENVOLVIMENTO')
    add_heading(doc, '3.1 Escopo e arquitetura da solução', 2)
    add_body(doc, 'O trabalho caracteriza-se como um desenvolvimento aplicado de software. Foram identificadas necessidades recorrentes do suporte a impressoras: visualizar equipamentos em um único painel, registrar informações de instalação, acompanhar toner e contadores, tratar alertas, registrar eventos e organizar materiais de reposição. A partir desses requisitos, foram definidos módulos de inventário, monitoramento, eventos e chamados, suprimentos, relatórios, clientes e configurações.')
    add_body(doc, 'O front-end utiliza React e Vite para a construção de uma interface responsiva. O back-end foi implementado em Node.js com Express, expondo endpoints REST para cadastro, edição, exclusão e consulta dos recursos do sistema. O SQLite foi adotado para persistir os registros durante a etapa de desenvolvimento. O modo --fake inicializa uma base de demonstração separada, permitindo testar as funcionalidades sem alterar os dados utilizados em uma instalação normal.')
    add_heading(doc, '3.2 Modelagem de dados', 2)
    add_body(doc, 'O modelo de dados é dividido em três grupos: clientes e sites; impressoras e históricos operacionais; e materiais e estoque. Um cliente pode ter vários sites e cada impressora pode ser associada a um cliente e a uma unidade. As leituras registram toner e contador de páginas; eventos e chamados mantêm o histórico técnico. No módulo de suprimentos, materiais possuem saldo, movimentações e compatibilidade com modelos de impressora.')
    add_figure(doc, 'original/image1.png', 'Figura 1 - Diagrama de Entidade Relacionamento do PrintWatch', 14.7)
    add_body(doc, 'A Figura 1 apresenta as principais entidades do banco de dados e suas relações. A organização permite associar ocorrências e leituras a cada impressora, além de vincular o consumo de materiais às rotinas de estoque.')
    add_figure(doc, 'original/image2.png', 'Figura 2 - Diagrama de classes do PrintWatch', 14.7)
    add_body(doc, 'O diagrama de classes da Figura 2 representa os objetos de domínio utilizados na aplicação. As operações de criação, listagem, atualização e exclusão correspondem aos endpoints oferecidos pela API REST.')

    add_heading(doc, '4 RESULTADOS E DISCUSSÃO')
    add_heading(doc, '4.1 Painel de monitoramento e alertas', 2)
    add_body(doc, 'O painel principal apresenta indicadores consolidados de impressoras cadastradas, equipamentos online, itens que exigem atenção e impressoras offline. A tabela detalha nome, setor, endereço IP, status, nível de toner, contador de páginas e data da última leitura. Os filtros permitem restringir a visualização por situação operacional e a busca facilita a localização de equipamentos por nome ou setor.')
    add_figure(doc, 'figura-1-painel.png', 'Figura 3 - Painel principal de monitoramento do PrintWatch', 15.0)
    add_body(doc, 'Na demonstração, o painel identifica uma impressora com toner abaixo do limite configurado e outra sem resposta de rede. As situações também são apresentadas no módulo de alertas automáticos, que informa a regra acionada e a gravidade do evento, conforme mostra a Figura 4.')
    add_figure(doc, 'figura-5-alertas.png', 'Figura 4 - Alertas automáticos de toner baixo e impressora offline', 14.8)

    add_heading(doc, '4.2 Detalhamento técnico e acompanhamento de chamados', 2)
    add_body(doc, 'Ao selecionar um equipamento, o sistema exibe seus dados operacionais e permite registrar eventos técnicos e abrir chamados de manutenção. O módulo registra descrição, prioridade, técnico responsável, status e solução aplicada. Também há uma estimativa de dias restantes de toner baseada nas leituras armazenadas, recurso que pode apoiar o planejamento de reposição.')
    add_figure(doc, 'figura-2-detalhes.png', 'Figura 5 - Detalhes da impressora, previsão de toner e abertura de chamado', 14.8)

    add_heading(doc, '4.3 Suprimentos, estoque e relatórios', 2)
    add_body(doc, 'O módulo de suprimentos permite cadastrar toner, cilindro ou peça de reposição, registrar estoque mínimo e lançar quantidades iniciais. As movimentações de entrada e saída são validadas para evitar saldo negativo. Esse controle é útil para relacionar o monitoramento de toner à disponibilidade dos materiais necessários ao atendimento.')
    add_figure(doc, 'figura-3-estoque.png', 'Figura 6 - Tela de cadastro de suprimentos e estoque', 14.8)
    add_body(doc, 'Os relatórios consolidam o total de impressoras, os contadores somados, os setores atendidos e as ocorrências registradas. A visualização por setor favorece a análise da distribuição do volume de impressão e pode apoiar decisões de manutenção e dimensionamento do parque.')
    add_figure(doc, 'figura-4-relatorios.png', 'Figura 7 - Relatório de uso por setor e histórico de ocorrências', 14.8)

    add_heading(doc, '5 CONCLUSÃO')
    add_body(doc, 'O desenvolvimento do PrintWatch resultou em uma aplicação web capaz de centralizar informações essenciais para o acompanhamento de impressoras em rede. A solução implementa cadastro de equipamentos, organização por cliente e site, visualização de status, contadores, níveis de toner, alertas, eventos técnicos, chamados, previsão de consumo e controle de estoque. O modo de demonstração permite avaliar os fluxos do sistema sem depender de uma impressora física conectada.')
    add_body(doc, 'Os resultados indicam que o sistema pode apoiar a identificação antecipada de situações como toner baixo e indisponibilidade de equipamento, contribuindo para uma atuação mais organizada do suporte técnico. Como continuidade, recomenda-se implementar um coletor SNMP para obter dados reais dos equipamentos, criar autenticação com níveis de permissão, disponibilizar notificações por canais de mensagem e integrar informações geográficas dos sites para visualização de ocorrências em mapa.')

    add_heading(doc, 'REFERÊNCIAS')
    references = [
        'EXPRESSJS. Express - Fast, unopinionated, minimalist web framework for Node.js. Disponível em: https://expressjs.com/. Acesso em: 20 out. 2026.',
        'IETF - INTERNET ENGINEERING TASK FORCE. RFC 3411: An Architecture for Describing Simple Network Management Protocol (SNMP) Management Frameworks. 2002. Disponível em: https://datatracker.ietf.org/doc/html/rfc3411. Acesso em: 20 out. 2026.',
        'KARDEC, Alan; NASCIF, Júlio. Manutenção: função estratégica. Rio de Janeiro: Qualitymark, 2013.',
        'NODE.JS. Node.js documentation. Disponível em: https://nodejs.org/docs/latest/api/. Acesso em: 20 out. 2026.',
        'PRESSMAN, Roger S.; MAXIM, Bruce R. Engenharia de software: uma abordagem profissional. 8. ed. Porto Alegre: AMGH, 2016.',
        'REACT. React documentation. Disponível em: https://react.dev/. Acesso em: 20 out. 2026.',
        'SQLITE. About SQLite. Disponível em: https://www.sqlite.org/about.html. Acesso em: 20 out. 2026.',
    ]
    for reference in references:
        add_reference(doc, reference)

    doc.core_properties.title = 'Desenvolvimento de um sistema web para monitoramento de impressoras em rede'
    doc.core_properties.author = 'João Pedro Lange'
    doc.save(OUTPUT)


if __name__ == '__main__':
    build()
